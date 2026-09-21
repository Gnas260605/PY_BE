import csv
import io
import os
from datetime import datetime
from io import StringIO

import openpyxl
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.db.connection import connection_scope
from app.reports import repository
from app.reports.schemas import TechnicianWorkloadResponse
from app.tickets.schemas import DashboardStatsResponse, TicketListQuery, TicketSummaryResponse


CATEGORY_LABELS = {
    "INCIDENT": "Sự cố",
    "SERVICE_REQUEST": "Yêu cầu dịch vụ",
    "MAINTENANCE": "Bảo trì",
    "ACCESS_REQUEST": "Cấp quyền",
}

PRIORITY_LABELS = {
    "URGENT": "Khẩn cấp",
    "HIGH": "Cao",
    "MEDIUM": "Trung bình",
    "LOW": "Thấp",
}

STATUS_LABELS = {
    "OPEN": "Chờ tiếp nhận",
    "ASSIGNED": "Đã phân công",
    "IN_PROGRESS": "Đang xử lý",
    "RESOLVED": "Đã giải quyết",
    "CLOSED": "Đã đóng",
}

CSV_HEADERS = [
    "Mã sự cố",
    "Tiêu đề",
    "Phân loại",
    "Mức độ ưu tiên",
    "Trạng thái",
    "Người yêu cầu",
    "KTV phụ trách",
    "Thiết bị liên quan",
    "Thời gian tạo",
    "Cập nhật lần cuối",
]


def get_technician_workload() -> list[TechnicianWorkloadResponse]:
    with connection_scope() as connection:
        rows = repository.get_technician_workload(connection)
    return [TechnicianWorkloadResponse(**row) for row in rows]


def export_tickets_csv(query: TicketListQuery) -> str:
    with connection_scope() as connection:
        rows = repository.list_tickets_for_export(
            connection,
            status=query.status,
            priority=query.priority,
            category=query.category,
            technician_id=query.technician_id,
            user_id=query.user_id,
            keyword=query.keyword,
        )

    output = StringIO()
    output.write("\ufeff")  # UTF-8 BOM for seamless Excel compatibility
    writer = csv.writer(output)
    writer.writerow(CSV_HEADERS)

    for row in rows:
        ticket_id = f"#TK-{row['id']:04d}" if row.get("id") else "-"
        category = CATEGORY_LABELS.get(row.get("category"), row.get("category") or "-")
        priority = PRIORITY_LABELS.get(row.get("priority"), row.get("priority") or "-")
        status_val = STATUS_LABELS.get(row.get("status"), row.get("status") or "-")
        creator = row.get("creator_name") or (f"User #{row.get('user_id')}" if row.get("user_id") else "-")
        technician = row.get("technician_name") or "Chưa phân công"

        if row.get("device_code") and row.get("device_name"):
            device = f"{row['device_code']} ({row['device_name']})"
        elif row.get("device_name"):
            device = str(row["device_name"])
        else:
            device = "-"

        created = str(row.get("created_at") or "-")
        updated = str(row.get("updated_at") or "-")

        writer.writerow([
            ticket_id,
            row.get("title") or "-",
            category,
            priority,
            status_val,
            creator,
            technician,
            device,
            created,
            updated,
        ])

    return output.getvalue()


def export_tickets_excel(query: TicketListQuery) -> io.BytesIO:
    with connection_scope() as connection:
        rows = repository.list_tickets_for_export(
            connection,
            status=query.status,
            priority=query.priority,
            category=query.category,
            technician_id=query.technician_id,
            user_id=query.user_id,
            keyword=query.keyword,
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tickets Export"

    for col_num, header_title in enumerate(CSV_HEADERS, 1):
        cell = ws.cell(row=1, column=col_num, value=header_title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_num, row in enumerate(rows, 2):
        ticket_id = f"#TK-{row['id']:04d}" if row.get("id") else "-"
        category = CATEGORY_LABELS.get(row.get("category"), row.get("category") or "-")
        priority = PRIORITY_LABELS.get(row.get("priority"), row.get("priority") or "-")
        status_val = STATUS_LABELS.get(row.get("status"), row.get("status") or "-")
        creator = row.get("creator_name") or (f"User #{row.get('user_id')}" if row.get("user_id") else "-")
        technician = row.get("technician_name") or "Chưa phân công"

        device = "-"
        if row.get("device_code") and row.get("device_name"):
            device = f"{row['device_code']} ({row['device_name']})"
        elif row.get("device_name"):
            device = str(row["device_name"])

        ws.cell(row=row_num, column=1, value=ticket_id)
        ws.cell(row=row_num, column=2, value=row.get("title") or "-")
        ws.cell(row=row_num, column=3, value=category)
        ws.cell(row=row_num, column=4, value=priority)
        ws.cell(row=row_num, column=5, value=status_val)
        ws.cell(row=row_num, column=6, value=creator)
        ws.cell(row=row_num, column=7, value=technician)
        ws.cell(row=row_num, column=8, value=device)
        ws.cell(row=row_num, column=9, value=str(row.get("created_at") or "-"))
        ws.cell(row=row_num, column=10, value=str(row.get("updated_at") or "-"))

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _register_unicode_fonts() -> tuple[str, str, str]:
    candidates = [
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf"),
        ("C:/Windows/Fonts/tahoma.ttf", "C:/Windows/Fonts/tahomabd.ttf", "C:/Windows/Fonts/tahoma.ttf"),
        ("C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/segoeuib.ttf", "C:/Windows/Fonts/segoeuii.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
    ]
    for reg, bold, italic in candidates:
        if os.path.exists(reg) and os.path.exists(bold):
            try:
                pdfmetrics.registerFont(TTFont("ReportUnicode", reg))
                pdfmetrics.registerFont(TTFont("ReportUnicode-Bold", bold))
                italic_path = italic if os.path.exists(italic) else reg
                pdfmetrics.registerFont(TTFont("ReportUnicode-Italic", italic_path))
                return "ReportUnicode", "ReportUnicode-Bold", "ReportUnicode-Italic"
            except Exception:
                pass
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"


def export_dashboard_pdf(stats: DashboardStatsResponse) -> io.BytesIO:
    font_reg, font_bold, font_italic = _register_unicode_fonts()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    elements = []

    # Typography & Styles
    style_org_left = ParagraphStyle(
        "OrgLeft",
        fontName=font_bold,
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#1e293b"),
        alignment=0,
    )
    style_org_sub = ParagraphStyle(
        "OrgSub",
        fontName=font_reg,
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#475569"),
        alignment=0,
    )
    style_national = ParagraphStyle(
        "NationalHeader",
        fontName=font_bold,
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#1e293b"),
        alignment=1,
    )
    style_motto = ParagraphStyle(
        "MottoHeader",
        fontName=font_bold,
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#334155"),
        alignment=1,
    )
    style_title = ParagraphStyle(
        "DocTitle",
        fontName=font_bold,
        fontSize=15,
        leading=19,
        textColor=colors.HexColor("#0f172a"),
        alignment=1,
    )
    style_subtitle = ParagraphStyle(
        "DocSubTitle",
        fontName=font_italic,
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#475569"),
        alignment=1,
    )
    style_meta = ParagraphStyle(
        "DocMeta",
        fontName=font_reg,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#334155"),
    )
    style_heading = ParagraphStyle(
        "SectionHeading",
        fontName=font_bold,
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#1e3a8a"),
        spaceBefore=10,
        spaceAfter=4,
    )
    style_cell = ParagraphStyle(
        "TableCell",
        fontName=font_reg,
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#1e293b"),
    )
    style_cell_center = ParagraphStyle(
        "TableCellCenter",
        fontName=font_reg,
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#1e293b"),
        alignment=1,
    )
    style_cell_bold = ParagraphStyle(
        "TableCellBold",
        fontName=font_bold,
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#0f172a"),
    )
    style_cell_bold_center = ParagraphStyle(
        "TableCellBoldCenter",
        fontName=font_bold,
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#0f172a"),
        alignment=1,
    )
    style_th = ParagraphStyle(
        "TableHeader",
        fontName=font_bold,
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=1,
    )

    # 1. Header Official Banner (2-Column Table)
    header_left = [
        Paragraph("HỆ THỐNG CS466 HELPDESK", style_org_left),
        Paragraph("TRUNG TÂM DỊCH VỤ & HỖ TRỢ KỸ THUẬT IT", style_org_sub),
        Paragraph("Mã số biểu mẫu: BM-IT-09/2026", style_org_sub),
    ]
    header_right = [
        Paragraph("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", style_national),
        Paragraph("Độc lập – Tự do – Hạnh phúc", style_motto),
        Paragraph("----------o0o----------", style_motto),
    ]

    header_table = Table([[header_left, header_right]], colWidths=[240, 280])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 14))

    # 2. Document Title Banner
    elements.append(Paragraph("BÁO CÁO THỐNG KÊ TỔNG QUAN HOẠT ĐỘNG HỖ TRỢ IT", style_title))
    elements.append(Paragraph("Phân tích dữ liệu vận hành hệ thống, tiến độ xử lý sự cố và tình trạng thiết bị", style_subtitle))
    elements.append(Spacer(1, 10))

    # 3. Report Metadata Strip
    now_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    meta_box = [
        [
            Paragraph(f"<b>Thời gian xuất báo cáo:</b> {now_str}", style_meta),
            Paragraph(f"<b>Phạm vi thống kê:</b> Toàn bộ hệ thống", style_meta),
        ],
        [
            Paragraph(f"<b>Đơn vị quản trị:</b> Ban Quản Trị Hệ Thống IT", style_meta),
            Paragraph(f"<b>Tổng số sự cố ghi nhận:</b> {stats.total_tickets} yêu cầu", style_meta),
        ]
    ]
    meta_table = Table(meta_box, colWidths=[260, 260])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor("#cbd5e1")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 12))

    # =========================================================================
    # Section 1: Tổng quan nguồn lực & hệ thống
    # =========================================================================
    elements.append(Paragraph("1. TỔNG QUAN TÀI NGUYÊN & VẬN HÀNH HỆ THỐNG", style_heading))
    
    status_counts = stats.status_counts or {}
    total_tickets = stats.total_tickets or 0
    resolved_total = status_counts.get("RESOLVED", 0) + status_counts.get("CLOSED", 0)
    in_progress_total = status_counts.get("OPEN", 0) + status_counts.get("ASSIGNED", 0) + status_counts.get("IN_PROGRESS", 0)
    completion_rate = f"{(resolved_total / max(1, total_tickets) * 100):.1f}%" if total_tickets > 0 else "100.0%"

    overview_rows = [
        [
            Paragraph("Hạng mục chỉ số", style_th),
            Paragraph("Số lượng", style_th),
            Paragraph("Đơn vị", style_th),
            Paragraph("Đánh giá & Diễn giải nghiệp vụ", style_th),
        ],
        [
            Paragraph("Tổng số sự cố & yêu cầu tiếp nhận", style_cell_bold),
            Paragraph(str(stats.total_tickets), style_cell_bold_center),
            Paragraph("Yêu cầu", style_cell_center),
            Paragraph("Toàn bộ khối lượng ticket phát sinh trên hệ thống", style_cell),
        ],
        [
            Paragraph("Sự cố đã khắc phục & Đóng hoàn tất", style_cell_bold),
            Paragraph(str(resolved_total), style_cell_bold_center),
            Paragraph("Yêu cầu", style_cell_center),
            Paragraph(f"Tỷ lệ giải quyết thành công đạt <b>{completion_rate}</b>", style_cell),
        ],
        [
            Paragraph("Sự cố đang trong tiến trình xử lý", style_cell_bold),
            Paragraph(str(in_progress_total), style_cell_bold_center),
            Paragraph("Yêu cầu", style_cell_center),
            Paragraph("Bao gồm trạng thái Chờ tiếp nhận, Đã gán & Đang sửa", style_cell),
        ],
        [
            Paragraph("Tổng số trang thiết bị IT quản lý", style_cell_bold),
            Paragraph(str(stats.total_devices), style_cell_bold_center),
            Paragraph("Thiết bị", style_cell_center),
            Paragraph("Máy tính trạm, máy in, thiết bị mạng văn phòng", style_cell),
        ],
        [
            Paragraph("Tổng số tài khoản người dùng hoạt động", style_cell_bold),
            Paragraph(str(stats.total_users), style_cell_bold_center),
            Paragraph("Tài khoản", style_cell_center),
            Paragraph("Người dùng nội bộ, Kỹ thuật viên & Quản trị viên", style_cell),
        ],
    ]

    tbl_overview = Table(overview_rows, colWidths=[180, 65, 55, 220])
    tbl_overview.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    elements.append(tbl_overview)
    elements.append(Spacer(1, 10))

    # =========================================================================
    # Section 2: Thống kê phân loại theo Trạng thái xử lý
    # =========================================================================
    elements.append(Paragraph("2. PHÂN BỔ THEO TIẾN ĐỘ & TRẠNG THÁI XỬ LÝ SỰ CỐ", style_heading))
    
    status_defs = [
        ("OPEN", "Chờ tiếp nhận (Open)", "Yêu cầu mới khởi tạo, chờ điều phối phân công"),
        ("ASSIGNED", "Đã phân công KTV (Assigned)", "Đã giao Kỹ thuật viên phụ trách, chuẩn bị xử lý"),
        ("IN_PROGRESS", "Đang tiến hành xử lý (In Progress)", "Kỹ thuật viên đang trực tiếp sửa chữa / cấu hình"),
        ("RESOLVED", "Đã giải quyết sự cố (Resolved)", "Đã khắc phục xong, chờ người dùng xác nhận"),
        ("CLOSED", "Đã đóng hoàn tất (Closed)", "Sự cố hoàn tất nghiệm thu và lưu trữ lịch sử"),
    ]

    status_table_rows = [
        [
            Paragraph("Trạng thái xử lý", style_th),
            Paragraph("Số lượng", style_th),
            Paragraph("Tỷ trọng", style_th),
            Paragraph("Định hướng & Quy trình xử lý", style_th),
        ]
    ]

    for code, label, note in status_defs:
        cnt = status_counts.get(code, 0)
        pct = f"{(cnt / max(1, total_tickets) * 100):.1f}%" if total_tickets > 0 else "0.0%"
        status_table_rows.append([
            Paragraph(label, style_cell_bold),
            Paragraph(str(cnt), style_cell_center),
            Paragraph(pct, style_cell_center),
            Paragraph(note, style_cell),
        ])

    status_table_rows.append([
        Paragraph("<b>TỔNG CỘNG</b>", style_cell_bold),
        Paragraph(f"<b>{total_tickets}</b>", style_cell_bold_center),
        Paragraph("<b>100.0%</b>", style_cell_bold_center),
        Paragraph("<b>Toàn bộ trạng thái sự cố</b>", style_cell),
    ])

    tbl_status = Table(status_table_rows, colWidths=[160, 60, 60, 240])
    tbl_status.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0284c7")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor("#f8fafc")]),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#e2e8f0")),
    ]))
    elements.append(tbl_status)
    elements.append(Spacer(1, 10))

    # =========================================================================
    # Section 3: Thống kê phân loại theo Mức độ ưu tiên
    # =========================================================================
    elements.append(Paragraph("3. PHÂN LOẠI THEO MỨC ĐỘ CẤP BÁCH & ƯU TIÊN (SLA)", style_heading))
    
    priority_counts = stats.priority_counts or {}
    priority_defs = [
        ("URGENT", "Khẩn cấp (Urgent)", "Thời gian đáp ứng < 1 giờ, giải quyết dứt điểm trong ngày"),
        ("HIGH", "Mức độ cao (High)", "Thời gian đáp ứng < 2 giờ, giải quyết trong vòng 4 - 8 giờ"),
        ("MEDIUM", "Mức trung bình (Medium)", "Xử lý theo hàng đợi tiêu chuẩn (trong vòng 24 giờ)"),
        ("LOW", "Mức độ thấp (Low)", "Xử lý theo kế hoạch công việc định kỳ (trong vòng 48 giờ)"),
    ]

    priority_table_rows = [
        [
            Paragraph("Mức độ ưu tiên", style_th),
            Paragraph("Số lượng", style_th),
            Paragraph("Tỷ trọng", style_th),
            Paragraph("Cam kết chất lượng dịch vụ (SLA)", style_th),
        ]
    ]

    for code, label, sla in priority_defs:
        cnt = priority_counts.get(code, 0)
        pct = f"{(cnt / max(1, total_tickets) * 100):.1f}%" if total_tickets > 0 else "0.0%"
        priority_table_rows.append([
            Paragraph(label, style_cell_bold),
            Paragraph(str(cnt), style_cell_center),
            Paragraph(pct, style_cell_center),
            Paragraph(sla, style_cell),
        ])

    priority_table_rows.append([
        Paragraph("<b>TỔNG CỘNG</b>", style_cell_bold),
        Paragraph(f"<b>{total_tickets}</b>", style_cell_bold_center),
        Paragraph("<b>100.0%</b>", style_cell_bold_center),
        Paragraph("<b>Toàn bộ mức ưu tiên sự cố</b>", style_cell),
    ])

    tbl_priority = Table(priority_table_rows, colWidths=[150, 60, 60, 250])
    tbl_priority.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#b91c1c")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor("#f8fafc")]),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor("#e2e8f0")),
    ]))
    elements.append(tbl_priority)
    elements.append(Spacer(1, 10))

    # =========================================================================
    # Section 4: Danh sách sự cố khẩn cấp tồn đọng (nếu có)
    # =========================================================================
    urgent_tickets = stats.urgent_tickets or []
    if urgent_tickets:
        elements.append(Paragraph("4. DANH SÁCH SỰ CỐ KHẨN CẤP ĐANG CẦN ƯU TIÊN XỬ LÝ", style_heading))
        urgent_rows = [
            [
                Paragraph("Mã", style_th),
                Paragraph("Tiêu đề sự cố", style_th),
                Paragraph("Ưu tiên", style_th),
                Paragraph("Trạng thái", style_th),
                Paragraph("Thời điểm tạo", style_th),
            ]
        ]
        for tck in urgent_tickets[:5]:
            t_id = f"#TK-{tck.id:04d}" if hasattr(tck, 'id') else f"#{tck.get('id', '-')}"
            t_title = tck.title if hasattr(tck, 'title') else tck.get('title', '-')
            t_prio = PRIORITY_LABELS.get(tck.priority if hasattr(tck, 'priority') else tck.get('priority'), "Khẩn")
            t_status = STATUS_LABELS.get(tck.status if hasattr(tck, 'status') else tck.get('status'), "Đang mở")
            t_created = str(tck.created_at if hasattr(tck, 'created_at') else tck.get('created_at', ''))[:16]
            urgent_rows.append([
                Paragraph(t_id, style_cell_bold_center),
                Paragraph(t_title, style_cell),
                Paragraph(t_prio, style_cell_center),
                Paragraph(t_status, style_cell_center),
                Paragraph(t_created, style_cell_center),
            ])

        tbl_urgent = Table(urgent_rows, colWidths=[65, 235, 65, 75, 80])
        tbl_urgent.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#475569")),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))
        elements.append(tbl_urgent)
        elements.append(Spacer(1, 10))

    # =========================================================================
    # Section 5: Chữ ký & Phê duyệt biểu mẫu
    # =========================================================================
    elements.append(Spacer(1, 15))
    sign_date_str = datetime.now().strftime("Ngày %d tháng %m năm %Y")
    
    sig_block = [
        [
            Paragraph("", style_cell_center),
            Paragraph(f"<i>TP. Hồ Chí Minh, {sign_date_str}</i>", style_cell_center),
        ],
        [
            Paragraph("<b>NGƯỜI LẬP BÁO CÁO</b><br/><i>(Ký và ghi rõ họ tên)</i>", style_cell_center),
            Paragraph("<b>TRƯỞNG BỘ PHẬN CNTT / QUẢN TRỊ VIÊN</b><br/><i>(Ký duyệt, đóng dấu xác nhận)</i>", style_cell_center),
        ],
        [
            Paragraph("<br/><br/><br/><b>Ban Vận Hành Hệ Thống IT</b>", style_cell_center),
            Paragraph("<br/><br/><br/><b>Quản Trị Viên Hệ Thống</b>", style_cell_center),
        ]
    ]
    sig_table = Table(sig_block, colWidths=[260, 260])
    sig_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
    ]))
    elements.append(sig_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer
